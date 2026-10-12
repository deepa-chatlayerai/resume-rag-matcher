"""Generate a synthetic resume + job-description dataset for the RAG matcher.

Outputs:
  resumes/<name>.(txt|docx|pdf)   36 resumes, section-structured, mixed formats
  jobs/<id>.txt                   6 job descriptions
  dataset_meta.json               ground-truth metadata for evaluation

Run:
  python generate_dataset.py

Requires python-docx (DOCX) and fpdf2 (PDF).
"""

from __future__ import annotations

import json
import os

import docx
from fpdf import FPDF
from fpdf.enums import XPos, YPos

RESUME_DIR = "resumes"
JOB_DIR = "jobs"

# role -> (title, core skills, extra skills, education field)
ROLES = {
    "backend": ("Backend Engineer",
                ["Python", "Django", "PostgreSQL", "REST APIs", "Redis"],
                ["Celery", "gRPC", "Microservices", "Docker"], "Computer Science"),
    "ml": ("Machine Learning Engineer",
           ["Python", "Machine Learning", "PyTorch", "scikit-learn", "NLP"],
           ["TensorFlow", "MLOps", "Pandas", "Spark"], "Machine Learning"),
    "data_science": ("Data Scientist",
                     ["Python", "Statistics", "Pandas", "SQL", "Machine Learning"],
                     ["R", "A/B Testing", "Tableau", "Experimentation"], "Statistics"),
    "frontend": ("Frontend Engineer",
                 ["JavaScript", "React", "TypeScript", "CSS", "HTML"],
                 ["Redux", "Next.js", "Accessibility", "Webpack"], "Computer Science"),
    "fullstack": ("Full Stack Developer",
                  ["JavaScript", "React", "Node.js", "TypeScript", "MongoDB"],
                  ["GraphQL", "Express", "AWS", "Docker"], "Software Engineering"),
    "devops": ("DevOps Engineer",
               ["AWS", "Terraform", "Docker", "Kubernetes", "CI/CD"],
               ["Ansible", "Python", "Go", "Prometheus"], "Computer Engineering"),
    "sre": ("Site Reliability Engineer",
            ["Kubernetes", "Prometheus", "Go", "Python", "Linux"],
            ["Grafana", "Terraform", "Incident Response", "SLOs"], "Computer Science"),
    "data_eng": ("Data Engineer",
                 ["Python", "Spark", "Airflow", "SQL", "ETL"],
                 ["Kafka", "Snowflake", "dbt", "AWS"], "Data Engineering"),
    "security": ("Security Engineer",
                 ["Python", "Threat Modeling", "Penetration Testing", "SIEM", "AWS"],
                 ["Burp Suite", "Go", "Cryptography", "SOC"], "Cybersecurity"),
    "mobile": ("Mobile Engineer",
               ["Swift", "Kotlin", "iOS", "Android", "Mobile"],
               ["React Native", "Flutter", "REST APIs", "CI/CD"], "Computer Science"),
    "qa": ("QA Automation Engineer",
           ["Python", "Selenium", "Test Automation", "CI/CD", "pytest"],
           ["Cypress", "Playwright", "API Testing", "Performance Testing"], "Information Technology"),
    "pm": ("Product Manager",
           ["Roadmapping", "Analytics", "SQL", "A/B Testing", "Agile"],
           ["Figma", "Stakeholder Management", "Go-to-Market", "Prioritization"], "Business Administration"),
}

SENIORITY = {
    "Junior": (2, "Junior", "Associate"),
    "Mid": (5, "", ""),
    "Senior": (9, "Senior", "Lead"),
}

FIRST = ["John", "Jane", "Raj", "Maria", "Liam", "Aisha", "Chen", "Sofia",
         "David", "Emma", "Omar", "Nina", "Carlos", "Priya", "Yuki", "Tariq",
         "Lena", "Diego", "Fatima", "Noah", "Grace", "Ivan", "Mei", "Hassan",
         "Zoe", "Arjun", "Clara", "Sven", "Leila", "Marco", "Anya", "Kofi",
         "Hana", "Pedro", "Sara", "Tom"]
LAST = ["Doe", "Smith", "Patel", "Garcia", "O'Connor", "Khan", "Wei", "Rossi",
        "Mueller", "Wilson", "Haddad", "Volkov", "Ruiz", "Nair", "Tanaka", "Aziz",
        "Berg", "Silva", "Noor", "Brown", "Lee", "Petrov", "Lin", "Ali",
        "Park", "Mehta", "Weiss", "Larsson", "Saab", "Conti", "Ivanova", "Mensah",
        "Kim", "Costa", "Cohen", "Fischer"]

COMPANIES = ["Acme Corp", "Beta Inc", "DataWorks", "CloudNine", "ShipFast",
             "VisionAI", "StreamHub", "BrightUI", "SafeGuard", "GrowthLabs",
             "Nodal", "PixelLabs", "CoreSystems", "Quantify", "NimbusIO"]

FMT_CYCLE = ["txt", "txt", "txt", "docx", "pdf"]  # ~60% txt, 20% docx, 20% pdf


def build_candidates() -> list[dict]:
    cands = []
    i = 0
    for role_key, (title, core, extra, edu) in ROLES.items():
        for level, (years, prefix, _alt) in SENIORITY.items():
            first, last = FIRST[i], LAST[i]
            name = f"{first} {last}"
            full_title = (f"{prefix} {title}").strip()
            skills = core + extra[: (1 if level == "Junior" else 2 if level == "Mid" else 3)]
            grad_year = 2026 - years - 4
            c = {
                "id": f"{role_key}_{level.lower()}_{i:02d}",
                "name": name,
                "role": role_key,
                "title": full_title,
                "seniority": level,
                "years": years,
                "skills": skills,
                "education": f"B.S. {edu}, {COMPANIES[i % len(COMPANIES)].split()[0]} University, {grad_year}",
                "company": COMPANIES[i % len(COMPANIES)],
                "company2": COMPANIES[(i + 5) % len(COMPANIES)],
                "fmt": FMT_CYCLE[i % len(FMT_CYCLE)],
            }
            cands.append(c)
            i += 1
    return cands


def resume_lines(c: dict) -> dict:
    """Return the structured sections of a resume as a dict of section->text."""
    start = 2026 - c["years"]
    prev_start = start - 2
    summary = (f"{c['title']} with {c['years']} years of experience. "
               f"Specializes in {', '.join(c['skills'][:3])}.")
    experience = [
        f"{c['title']}, {c['company']} ({start}-present): "
        f"delivered projects using {', '.join(c['skills'][:3])}.",
        f"{c['title'].replace('Senior ', '').replace('Lead ', '')}, {c['company2']} "
        f"({prev_start}-{start}): built systems with {', '.join(c['skills'][3:5] or c['skills'][:2])}.",
    ]
    return {
        "Summary": summary,
        "Skills": ", ".join(c["skills"]),
        "Experience": experience,
        "Education": c["education"],
    }


def render_text(c: dict, s: dict) -> str:
    lines = [c["name"], c["title"], f"{c['name'].split()[0].lower()}.{c['name'].split()[1].lower().strip(chr(39))}@example.com",
             "", "SUMMARY", s["Summary"], "", "SKILLS", s["Skills"], "", "EXPERIENCE"]
    lines += [f"- {e}" for e in s["Experience"]]
    lines += ["", "EDUCATION", s["Education"], ""]
    return "\n".join(lines)


def write_txt(path, c, s):
    with open(path, "w", encoding="utf-8") as f:
        f.write(render_text(c, s))


def write_docx(path, c, s):
    d = docx.Document()
    d.add_heading(c["name"], 0)
    d.add_paragraph(c["title"])
    d.add_heading("Summary", 1); d.add_paragraph(s["Summary"])
    d.add_heading("Skills", 1); d.add_paragraph(s["Skills"])
    d.add_heading("Experience", 1)
    for e in s["Experience"]:
        d.add_paragraph(e, style="List Bullet")
    d.add_heading("Education", 1); d.add_paragraph(s["Education"])
    d.save(path)


def write_pdf(path, c, s):
    nl = {"new_x": XPos.LMARGIN, "new_y": YPos.NEXT}
    pdf = FPDF(); pdf.add_page()
    pdf.set_font("Helvetica", "B", 16); pdf.cell(0, 10, c["name"], **nl)
    pdf.set_font("Helvetica", "", 12); pdf.cell(0, 8, c["title"], **nl); pdf.ln(3)

    def section(title, body):
        pdf.set_font("Helvetica", "B", 13); pdf.cell(0, 8, title, **nl)
        pdf.set_font("Helvetica", "", 11)
        for line in body:
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(pdf.epw, 6, line.encode("latin-1", "replace").decode("latin-1"),
                           new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(2)

    section("Summary", [s["Summary"]])
    section("Skills", [s["Skills"]])
    section("Experience", [f"- {e}" for e in s["Experience"]])
    section("Education", [s["Education"]])
    pdf.output(path)


WRITERS = {"txt": write_txt, "docx": write_docx, "pdf": write_pdf}

# Job descriptions: target_roles + must-have skills + min years drive evaluation.
JOBS = [
    {"id": "jd_ml_engineer", "title": "Senior Machine Learning Engineer",
     "target_roles": ["ml", "data_science"], "must_have": ["Python", "Machine Learning"],
     "min_years": 5,
     "text": ("We are hiring a Senior Machine Learning Engineer to build and deploy "
              "ML models in production. Must have 5+ years of Python and hands-on "
              "machine learning experience (PyTorch or TensorFlow), including NLP. "
              "Experience with MLOps and large-scale data pipelines is a plus.")},
    {"id": "jd_backend_python", "title": "Backend Engineer (Python)",
     "target_roles": ["backend", "data_eng"], "must_have": ["Python", "REST APIs"],
     "min_years": 3,
     "text": ("Backend Engineer to build scalable REST APIs and services. Strong "
              "Python and relational database (PostgreSQL) skills required, 3+ years. "
              "Familiarity with Redis, Docker, and microservices preferred.")},
    {"id": "jd_frontend_react", "title": "Frontend Engineer (React)",
     "target_roles": ["frontend", "fullstack"], "must_have": ["React", "JavaScript"],
     "min_years": 3,
     "text": ("Frontend Engineer with strong React and JavaScript/TypeScript skills, "
              "3+ years. You will build accessible, performant UIs. Experience with "
              "Next.js and design systems is a plus.")},
    {"id": "jd_devops_sre", "title": "DevOps / SRE",
     "target_roles": ["devops", "sre"], "must_have": ["Kubernetes", "Terraform"],
     "min_years": 4,
     "text": ("DevOps/SRE to own our cloud infrastructure and reliability. Must have "
              "Kubernetes, Terraform, and AWS experience, 4+ years. CI/CD, Prometheus, "
              "and incident response experience valued.")},
    {"id": "jd_data_engineer", "title": "Data Engineer",
     "target_roles": ["data_eng"], "must_have": ["Spark", "Airflow", "SQL"],
     "min_years": 4,
     "text": ("Data Engineer to build batch and streaming pipelines. Must have Spark, "
              "Airflow, and SQL, 4+ years. Kafka, Snowflake, and dbt experience a plus.")},
    {"id": "jd_security", "title": "Security Engineer",
     "target_roles": ["security"], "must_have": ["Threat Modeling", "Penetration Testing"],
     "min_years": 4,
     "text": ("Security Engineer to lead application security. Must have threat "
              "modeling and penetration testing experience, 4+ years. Python and "
              "cloud (AWS) security knowledge required.")},
]


def main():
    os.makedirs(RESUME_DIR, exist_ok=True)
    os.makedirs(JOB_DIR, exist_ok=True)
    candidates = build_candidates()
    meta = {"resumes": {}, "jobs": {}}

    for c in candidates:
        s = resume_lines(c)
        fname = f"{c['name'].replace(' ', '_').replace(chr(39), '')}.{c['fmt']}"
        path = os.path.join(RESUME_DIR, fname)
        WRITERS[c["fmt"]](path, c, s)
        meta["resumes"][fname] = {
            "name": c["name"], "role": c["role"], "title": c["title"],
            "seniority": c["seniority"], "years": c["years"], "skills": c["skills"],
            "education": c["education"],
        }

    for j in JOBS:
        with open(os.path.join(JOB_DIR, f"{j['id']}.txt"), "w", encoding="utf-8") as f:
            f.write(f"{j['title']}\n\n{j['text']}\n")
        meta["jobs"][j["id"]] = j

    with open("dataset_meta.json", "w", encoding="utf-8") as f:
        json.dump(meta, f, indent=2)

    fmt_counts = {}
    for c in candidates:
        fmt_counts[c["fmt"]] = fmt_counts.get(c["fmt"], 0) + 1
    print(f"Generated {len(candidates)} resumes {fmt_counts} and {len(JOBS)} job descriptions.")
    print("Wrote dataset_meta.json")


if __name__ == "__main__":
    main()
